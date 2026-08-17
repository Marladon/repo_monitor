import urllib.request
import urllib.parse
import json
import ssl
import io
import os
import re
from datetime import datetime, timezone
from openpyxl import load_workbook

# ── CONFIG ────────────────────────────────────────────────
from pathlib import Path

def load_env():
    env_file = Path(__file__).parent / ".env"
    for line in env_file.read_text().splitlines():
        if "=" in line and not line.startswith("#"):
            key, val = line.split("=", 1)
            os.environ[key.strip()] = val.strip()

load_env()

SHEET_ID     = os.environ["SHEET_ID"]
GITLAB_TOKEN = os.environ["GITLAB_TOKEN"]
GITHUB_TOKEN = os.environ["GITHUB_TOKEN"]

LAST_CHECK_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "last_check.txt")
# ── END CONFIG ────────────────────────────────────────────

# ── Исключения слияний ────────────────────────────────────
# Загружаются из exceptions.json рядом со скриптом
def load_exceptions():
    path = Path(__file__).parent / "exceptions.json"
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    return {}

MERGE_EXCEPTIONS = load_exceptions()
# ── END MERGE_EXCEPTIONS ──────────────────────────────────


ctx = ssl._create_unverified_context()


def api_get(url, headers):
    req = urllib.request.Request(url, headers=headers)
    with urllib.request.urlopen(req, timeout=10, context=ctx) as resp:
        return json.loads(resp.read())


def github_headers():
    return {
        "Authorization": f"token {GITHUB_TOKEN}",
        "Accept": "application/vnd.github.v3+json"
    }


def gitlab_headers():
    return {"PRIVATE-TOKEN": GITLAB_TOKEN}


def parse_repo(url):
    parts = url.rstrip("/").split("/")
    return parts[-2], parts[-1]


# ── Ветки ─────────────────────────────────────────────────

def get_branches(url):
    try:
        namespace, repo = parse_repo(url)
        if "github.com" in url:
            data = api_get(
                f"https://api.github.com/repos/{namespace}/{repo}/branches?per_page=100",
                github_headers()
            )
            return [b["name"] for b in data]
        else:
            path = urllib.parse.quote(f"{namespace}/{repo}", safe="")
            data = api_get(
                f"https://gitlab.ximc.ru/api/v4/projects/{path}/repository/branches?per_page=100",
                gitlab_headers()
            )
            return [b["name"] for b in data]
    except Exception:
        return None


def get_dev_branches(branches):
    if branches is None:
        return None
    return [b for b in branches if b.startswith("dev-") or b == "dev"]


def sort_dev_branches(dev_branches):
    """Сортировка dev-веток по версии: dev-1.4 < dev-1.5 < dev-2.0"""
    def version_key(name):
        nums = re.findall(r'\d+', name)
        return [int(n) for n in nums]
    return sorted(dev_branches, key=version_key)




def is_merge_excluded(repo_name, from_branch, to_branch):
    """Проверяет попадает ли пара веток в список исключений."""
    exceptions = MERGE_EXCEPTIONS.get(repo_name, [])
    for exc_from, exc_to in exceptions:
        from_match = exc_from == "*" or exc_from == from_branch
        to_match = exc_to == "*" or exc_to == to_branch
        if from_match and to_match:
            return True
    return False

# ── Коммиты с даты ────────────────────────────────────────

def get_branches_with_new_commits(url, since_dt):
    since_str = since_dt.strftime("%Y-%m-%dT%H:%M:%SZ")
    namespace, repo = parse_repo(url)
    changed = []

    try:
        branches = get_branches(url)
        if not branches:
            return None

        for branch in branches:
            if "github.com" in url:
                data = api_get(
                    f"https://api.github.com/repos/{namespace}/{repo}/commits"
                    f"?sha={urllib.parse.quote(branch)}&since={since_str}&per_page=1",
                    github_headers()
                )
                if data:
                    changed.append(branch)
            else:
                path = urllib.parse.quote(f"{namespace}/{repo}", safe="")
                data = api_get(
                    f"https://gitlab.ximc.ru/api/v4/projects/{path}/repository/commits"
                    f"?ref_name={urllib.parse.quote(branch)}&since={since_str}&per_page=1",
                    gitlab_headers()
                )
                if data:
                    changed.append(branch)

        return changed

    except Exception:
        return None


# ── Проверка слияния ──────────────────────────────────────

def check_merge(url, from_branch, to_branch):
    """Проверяет есть ли в from_branch коммиты которых нет в to_branch.
    Возвращает True если всё влито, False если нет, None если ошибка."""
    try:
        namespace, repo = parse_repo(url)
        if "github.com" in url:
            # compare base=to_branch, head=from_branch
            # ahead_by > 0 означает что from_branch содержит коммиты которых нет в to_branch
            data = api_get(
                f"https://api.github.com/repos/{namespace}/{repo}/compare"
                f"/{urllib.parse.quote(to_branch)}...{urllib.parse.quote(from_branch)}",
                github_headers()
            )
            return data.get("ahead_by", 1) == 0
        else:
            path = urllib.parse.quote(f"{namespace}/{repo}", safe="")
            data = api_get(
                f"https://gitlab.ximc.ru/api/v4/projects/{path}/repository/compare"
                f"?from={urllib.parse.quote(to_branch)}&to={urllib.parse.quote(from_branch)}",
                gitlab_headers()
            )
            commits = data.get("commits", [])
            return len(commits) == 0
    except Exception:
        return None


# ── Утилиты ───────────────────────────────────────────────

def load_last_check():
    if os.path.exists(LAST_CHECK_FILE):
        with open(LAST_CHECK_FILE) as f:
            ts = f.read().strip()
            return datetime.strptime(ts, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=timezone.utc)
    return None


def save_last_check(dt):
    with open(LAST_CHECK_FILE, "w") as f:
        f.write(dt.strftime("%Y-%m-%dT%H:%M:%SZ"))


def load_sheet():
    url = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=15, context=ctx) as resp:
        return load_workbook(io.BytesIO(resp.read()))


# ── MAIN ──────────────────────────────────────────────────

now = datetime.now(timezone.utc)
last_check = load_last_check()

print("Loading spreadsheet...")
wb = load_sheet()
ws = wb.active

repos = []
for row in ws.iter_rows(min_row=2):
    cell = row[0]
    if cell.value and cell.hyperlink:
        repos.append({"name": cell.value, "url": cell.hyperlink.target})

print(f"Total repos: {len(repos)}")

# ── Блок 1: dev-ветки ─────────────────────────────────────
print("\n" + "=" * 50)
print("БЛОК 1 — Репозитории с несколькими dev-ветками")
print("=" * 50)

multi_dev = []
errors_dev = []

for repo in repos:
    name, url = repo["name"], repo["url"]
    print(f"  Checking {name}...", end="\r")

    branches = get_branches(url)
    dev = get_dev_branches(branches)

    if dev is None:
        errors_dev.append(name)
    elif len(dev) > 1:
        sorted_dev = sort_dev_branches(dev)
        multi_dev.append({"name": name, "url": url, "dev_branches": sorted_dev})

print(" " * 50)

if multi_dev:
    for r in multi_dev:
        print(f"  MULTIPLE  {r['name']} — {len(r['dev_branches'])} dev ветки: {', '.join(r['dev_branches'])}")
else:
    print("  Репозиториев с несколькими dev-ветками не найдено")

if errors_dev:
    print(f"\n  Ошибки API ({len(errors_dev)}): {', '.join(errors_dev)}")

print(f"\nАнализ завершён. Репозиториев с dev-ветками > 1: {len(multi_dev)}")

# ── Блок 2: новые коммиты + проверка слияний ──────────────
print("\n" + "=" * 50)
print("БЛОК 2 — Изменения с последней проверки")
print("=" * 50)

if last_check is None:
    print("  Первый запуск — сохраняем время, при следующем запуске появятся изменения.")
    save_last_check(now)
else:
    print(f"  Последняя проверка: {last_check.strftime('%Y-%m-%d %H:%M UTC')}\n")

    changed_repos = []
    errors_commits = []

    for repo in multi_dev:
        name, url = repo["name"], repo["url"]
        dev_branches = repo["dev_branches"]
        print(f"  Checking {name}...", end="\r")

        changed_branches = get_branches_with_new_commits(url, last_check)

        if changed_branches is None:
            errors_commits.append(name)
        elif changed_branches:
            # Проверяем слияния между соседними dev-ветками
            merge_status = []
            for i in range(len(dev_branches) - 1):
                from_b = dev_branches[i]
                to_b = dev_branches[i + 1]
                if is_merge_excluded(name, from_b, to_b):
                    merge_status.append(f"  В репозитории {name} ветка {from_b} -> {to_b}: исключено")
                    continue
                merged = check_merge(url, from_b, to_b)
                if merged is None:
                    merge_status.append(f"  В репозитории {name} ветка {from_b} -> {to_b}: ошибка проверки")
                elif merged:
                    pass  # влито — не выводим
                else:
                    merge_status.append(f"  В репозитории {name} ветка {from_b} не влита в {to_b}")

            changed_repos.append({
                "name": name,
                "branches": changed_branches,
                "merge_status": merge_status
            })

    print(" " * 50)

    if changed_repos:
        for r in changed_repos:
            print(f"  {r['name']} — изменения в: {', '.join(r['branches'])}")
            for line in r["merge_status"]:
                print(f"    {line}")
            print()
    else:
        print("  Изменений не найдено")

    if errors_commits:
        print(f"\n  Ошибки API ({len(errors_commits)}): {', '.join(errors_commits)}")

    save_last_check(now)
    print(f"\nПроверка завершена. {now.strftime('%Y-%m-%d %H:%M UTC')} сохранено как время последней проверки.")