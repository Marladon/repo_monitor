import urllib.request
import urllib.parse
import json
import ssl
import io
import os
import re
from datetime import datetime, timezone
from openpyxl import load_workbook

# ── CONFIG ────────────────────────────────────────────────
from pathlib import Path

def load_env():
    env_file = Path(__file__).parent / ".env"
    for line in env_file.read_text().splitlines():
        if "=" in line and not line.startswith("#"):
            key, val = line.split("=", 1)
            os.environ[key.strip()] = val.strip()

load_env()

SHEET_ID     = os.environ["SHEET_ID"]
GITLAB_TOKEN = os.environ["GITLAB_TOKEN"]
GITHUB_TOKEN = os.environ["GITHUB_TOKEN"]
TG_TOKEN     = os.environ["TG_TOKEN"]
TG_USER_ID   = int(os.environ["TG_USER_ID"])

LAST_CHECK_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "last_check.txt")
# ── END CONFIG ────────────────────────────────────────────

ctx = ssl._create_unverified_context()


# ── Telegram API ──────────────────────────────────────────

def tg_send(chat_id, text):
    url = f"https://api.telegram.org/bot{TG_TOKEN}/sendMessage"
    data = json.dumps({
        "chat_id": chat_id,
        "text": text,
        "parse_mode": "HTML"
    }).encode()
    req = urllib.request.Request(url, data=data, headers={"Content-Type": "application/json"})
    urllib.request.urlopen(req, timeout=10, context=ctx)


def tg_send_keyboard(chat_id, text):
    url = f"https://api.telegram.org/bot{TG_TOKEN}/sendMessage"
    data = json.dumps({
        "chat_id": chat_id,
        "text": text,
        "parse_mode": "HTML",
        "reply_markup": {
            "inline_keyboard": [[
                {"text": "🔍 Запустить проверку", "callback_data": "run"}
            ]]
        }
    }).encode()
    req = urllib.request.Request(url, data=data, headers={"Content-Type": "application/json"})
    try:
        with urllib.request.urlopen(req, timeout=10, context=ctx) as resp:
            pass
    except urllib.error.HTTPError as e:
        print("TG Error:", e.read().decode())
        raise


def tg_answer_callback(callback_id):
    url = f"https://api.telegram.org/bot{TG_TOKEN}/answerCallbackQuery"
    data = json.dumps({"callback_query_id": callback_id}).encode()
    req = urllib.request.Request(url, data=data, headers={"Content-Type": "application/json"})
    urllib.request.urlopen(req, timeout=10, context=ctx)


def tg_get_updates(offset=None):
    url = f"https://api.telegram.org/bot{TG_TOKEN}/getUpdates?timeout=30"
    if offset:
        url += f"&offset={offset}"
    req = urllib.request.Request(url)
    with urllib.request.urlopen(req, timeout=35, context=ctx) as resp:
        return json.loads(resp.read())


# ── Git API ───────────────────────────────────────────────

def api_get(url, headers):
    req = urllib.request.Request(url, headers=headers)
    with urllib.request.urlopen(req, timeout=10, context=ctx) as resp:
        return json.loads(resp.read())


def github_headers():
    return {"Authorization": f"token {GITHUB_TOKEN}", "Accept": "application/vnd.github.v3+json"}


def gitlab_headers():
    return {"PRIVATE-TOKEN": GITLAB_TOKEN}


def parse_repo(url):
    parts = url.rstrip("/").split("/")
    return parts[-2], parts[-1]


def get_branches(url):
    try:
        namespace, repo = parse_repo(url)
        if "github.com" in url:
            data = api_get(
                f"https://api.github.com/repos/{namespace}/{repo}/branches?per_page=100",
                github_headers()
            )
            return [b["name"] for b in data]
        else:
            path = urllib.parse.quote(f"{namespace}/{repo}", safe="")
            data = api_get(
                f"https://gitlab.ximc.ru/api/v4/projects/{path}/repository/branches?per_page=100",
                gitlab_headers()
            )
            return [b["name"] for b in data]
    except Exception:
        return None


def get_dev_branches(branches):
    if branches is None:
        return None
    return [b for b in branches if b.startswith("dev-") or b == "dev"]


def sort_dev_branches(dev_branches):
    def version_key(name):
        nums = re.findall(r'\d+', name)
        return [int(n) for n in nums]
    return sorted(dev_branches, key=version_key)


def get_branches_with_new_commits(url, since_dt):
    since_str = since_dt.strftime("%Y-%m-%dT%H:%M:%SZ")
    namespace, repo = parse_repo(url)
    changed = []
    try:
        branches = get_branches(url)
        if not branches:
            return None
        for branch in branches:
            if "github.com" in url:
                data = api_get(
                    f"https://api.github.com/repos/{namespace}/{repo}/commits"
                    f"?sha={urllib.parse.quote(branch)}&since={since_str}&per_page=1",
                    github_headers()
                )
                if data:
                    changed.append(branch)
            else:
                path = urllib.parse.quote(f"{namespace}/{repo}", safe="")
                data = api_get(
                    f"https://gitlab.ximc.ru/api/v4/projects/{path}/repository/commits"
                    f"?ref_name={urllib.parse.quote(branch)}&since={since_str}&per_page=1",
                    gitlab_headers()
                )
                if data:
                    changed.append(branch)
        return changed
    except Exception:
        return None


def check_merge(url, from_branch, to_branch):
    try:
        namespace, repo = parse_repo(url)
        if "github.com" in url:
            data = api_get(
                f"https://api.github.com/repos/{namespace}/{repo}/compare"
                f"/{urllib.parse.quote(to_branch)}...{urllib.parse.quote(from_branch)}",
                github_headers()
            )
            return data.get("ahead_by", 1) == 0
        else:
            path = urllib.parse.quote(f"{namespace}/{repo}", safe="")
            data = api_get(
                f"https://gitlab.ximc.ru/api/v4/projects/{path}/repository/compare"
                f"?from={urllib.parse.quote(to_branch)}&to={urllib.parse.quote(from_branch)}",
                gitlab_headers()
            )
            return len(data.get("commits", [])) == 0
    except Exception:
        return None


# ── Утилиты ───────────────────────────────────────────────

def load_last_check():
    if os.path.exists(LAST_CHECK_FILE):
        with open(LAST_CHECK_FILE) as f:
            ts = f.read().strip()
            return datetime.strptime(ts, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=timezone.utc)
    return None


def save_last_check(dt):
    with open(LAST_CHECK_FILE, "w") as f:
        f.write(dt.strftime("%Y-%m-%dT%H:%M:%SZ"))


def load_sheet():
    url = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=15, context=ctx) as resp:
        return load_workbook(io.BytesIO(resp.read()))


# ── Основная проверка ─────────────────────────────────────

def run_check(chat_id):
    now = datetime.now(timezone.utc)
    last_check = load_last_check()

    tg_send(chat_id, "⏳ Загружаю таблицу репозиториев...")

    wb = load_sheet()
    ws = wb.active

    repos = []
    for row in ws.iter_rows(min_row=2):
        cell = row[0]
        if cell.value and cell.hyperlink:
            repos.append({"name": cell.value, "url": cell.hyperlink.target})

    tg_send(chat_id, f"📋 Найдено репозиториев: <b>{len(repos)}</b>\nПроверяю dev-ветки...")

    # Блок 1
    multi_dev = []
    errors_dev = []

    for repo in repos:
        name, url = repo["name"], repo["url"]
        branches = get_branches(url)
        dev = get_dev_branches(branches)
        if dev is None:
            errors_dev.append(name)
        elif len(dev) > 1:
            multi_dev.append({"name": name, "url": url, "dev_branches": sort_dev_branches(dev)})

    msg1 = "<b>📊 БЛОК 1 — Репозитории с несколькими dev-ветками</b>\n"
    msg1 += "<i>Репозитории в которых нужно следить за подлитием</i>\n"
    msg1 += "—" * 20 + "\n\n"
    if multi_dev:
        for r in multi_dev:
            msg1 += f"📁 <b>{r['name']}</b> ({len(r['dev_branches'])} ветки):\n"
            msg1 += f"   <code>{', '.join(r['dev_branches'])}</code>\n\n"
    else:
        msg1 += "Репозиториев с несколькими dev-ветками не найдено\n"
    if errors_dev:
        msg1 += f"\n⚠️ Ошибки API ({len(errors_dev)}): {', '.join(errors_dev)}"
    msg1 += f"—" * 20 + "\n"
    msg1 += f"Всего: <b>{len(multi_dev)}</b> репозиториев"

    tg_send(chat_id, msg1)

    # Блок 2
    if last_check is None:
        tg_send(chat_id, "⏰ <b>БЛОК 2 — Изменения</b>\n\nПервый запуск — время сохранено.\nПри следующем запуске появятся изменения.")
        save_last_check(now)
        return

    tg_send(chat_id, f"🔍 Проверяю изменения с <b>{last_check.strftime('%Y-%m-%d %H:%M UTC')}</b>...")

    changed_repos = []
    errors_commits = []

    for repo in multi_dev:
        name, url = repo["name"], repo["url"]
        dev_branches = repo["dev_branches"]

        changed_branches = get_branches_with_new_commits(url, last_check)

        if changed_branches is None:
            errors_commits.append(name)
        elif changed_branches:
            merge_status = []
            for i in range(len(dev_branches) - 1):
                from_b = dev_branches[i]
                to_b = dev_branches[i + 1]
                merged = check_merge(url, from_b, to_b)
                if merged is None:
                    merge_status.append(f"   <code>{from_b}</code> → <code>{to_b}</code>: ❓ ошибка проверки")
                elif merged:
                    merge_status.append(f"   <code>{from_b}</code> → <code>{to_b}</code>: ✅ влито")
                else:
                    merge_status.append(f"   <code>{from_b}</code> → <code>{to_b}</code>: ❌ не влито")

            changed_repos.append({"name": name, "branches": changed_branches, "merge_status": merge_status})

    msg2 = f"<b>🔄 БЛОК 2 — Изменения с {last_check.strftime('%Y-%m-%d %H:%M UTC')}</b>\n"
    msg2 += "—" * 20 + "\n\n"
    if changed_repos:
        for r in changed_repos:
            msg2 += f"📦 <b>{r['name']}</b>\n"
            msg2 += f"   Ветки с изменениями: <code>{', '.join(r['branches'])}</code>\n"
            msg2 += "\n".join(r["merge_status"]) + "\n\n"
    else:
        msg2 += "Изменений не найдено"

    if errors_commits:
        msg2 += f"\n⚠️ Ошибки API ({len(errors_commits)}): {', '.join(errors_commits)}"

    tg_send(chat_id, msg2)
    save_last_check(now)
    tg_send_keyboard(chat_id, f"✅ Проверка завершена\n<i>{now.strftime('%Y-%m-%d %H:%M UTC')}</i>")


# ── Polling loop ──────────────────────────────────────────

def main():
    print("Бот запущен. Ожидаю команды...")
    tg_send_keyboard(TG_USER_ID, "👋 Бот запущен!\nНажми кнопку чтобы запустить проверку репозиториев.")

    offset = None
    while True:
        try:
            updates = tg_get_updates(offset)
            for update in updates.get("result", []):
                offset = update["update_id"] + 1

                if "message" in update:
                    msg = update["message"]
                    if msg.get("from", {}).get("id") != TG_USER_ID:
                        continue
                    if msg.get("text") == "/start":
                        tg_send_keyboard(TG_USER_ID, "👋 Привет!\nНажми кнопку чтобы запустить проверку.")

                elif "callback_query" in update:
                    cb = update["callback_query"]
                    if cb.get("from", {}).get("id") != TG_USER_ID:
                        continue
                    if cb.get("data") == "run":
                        tg_answer_callback(cb["id"])
                        try:
                            run_check(TG_USER_ID)
                        except Exception as e:
                            tg_send(TG_USER_ID, f"❌ Ошибка во время проверки:\n<code>{e}</code>")

        except KeyboardInterrupt:
            print("Бот остановлен.")
            break
        except Exception as e:
            print(f"Ошибка polling: {e}")


if __name__ == "__main__":
    main()