import urllib.request
import urllib.parse
import json
import ssl
import io
import os
import re
from openpyxl import load_workbook
from pathlib import Path

# ENV
def load_env():
    env_file = Path(__file__).parent / ".env"
    for line in env_file.read_text().splitlines():
        if "=" in line and not line.startswith("#"):
            key, val = line.split("=", 1)
            os.environ[key.strip()] = val.strip()

load_env()

SHEET_ID = os.environ["SHEET_ID"]
GITLAB_TOKEN = os.environ["GITLAB_TOKEN"]
GITHUB_TOKEN = os.environ["GITHUB_TOKEN"]

ctx = ssl._create_unverified_context()

# API
def api_get(url, headers):
    req = urllib.request.Request(url, headers=headers)
    with urllib.request.urlopen(req, timeout=10, context=ctx) as resp:
        return json.loads(resp.read())


def github_headers():
    return {
        "Authorization": f"token {GITHUB_TOKEN}",
        "Accept": "application/vnd.github.v3+json"
    }


def gitlab_headers():
    return {"PRIVATE-TOKEN": GITLAB_TOKEN}


def parse_repo(url):
    parts = url.rstrip("/").split("/")
    return parts[-2], parts[-1]


# SHEET
def load_sheet():
    url = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=15, context=ctx) as resp:
        return load_workbook(io.BytesIO(resp.read()))


# BRANCHES
def get_branches(url):
    try:
        namespace, repo = parse_repo(url)

        if "github.com" in url:
            data = api_get(
                f"https://api.github.com/repos/{namespace}/{repo}/branches?per_page=100",
                github_headers()
            )
            return [b["name"] for b in data]

        else:
            path = urllib.parse.quote(f"{namespace}/{repo}", safe="")
            data = api_get(
                f"https://gitlab.ximc.ru/api/v4/projects/{path}/repository/branches?per_page=100",
                gitlab_headers()
            )
            return [b["name"] for b in data]

    except Exception:
        return None


# T-BRANCH FILTER
T_BRANCH_PATTERN = re.compile(r"^t-\d+.*$")


def get_t_branches(branches):
    if branches is None:
        return None
    return [b for b in branches if T_BRANCH_PATTERN.match(b)]


# MAIN
print("Loading spreadsheet...")

wb = load_sheet()
ws = wb.active

repos = []
for row in ws.iter_rows(min_row=2):
    cell = row[0]
    if cell.value and cell.hyperlink:
        repos.append({
            "name": cell.value,
            "url": cell.hyperlink.target
        })

print(f"Total repos: {len(repos)}")

print("\n" + "=" * 50)
print("Репозитории с t-ветками")
print("=" * 50)

found = []
errors = []

for repo in repos:
    name, url = repo["name"], repo["url"]
    print(f"Checking {name}...", end="\r")

    branches = get_branches(url)
    t_branches = get_t_branches(branches)

    if t_branches is None:
        errors.append(name)
    elif t_branches:
        found.append((name, sorted(t_branches)))

print(" " * 50)

if found:
    for name, branches in found:
        print(f"{name}:")
        for b in branches:
            print(f"  - {b}")
        print()
else:
    print("Репозиториев с t-ветками не найдено")

if errors:
    print(f"\nОшибки API ({len(errors)}):")
    for e in errors:
        print(f"  - {e}")

print(f"\nГотово. Найдено репозиториев: {len(found)}")